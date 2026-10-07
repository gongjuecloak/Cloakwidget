# -*- coding: utf-8 -*-
"""批量多文件验证：同一次请求里传多个源文件，检查合并后的产出与单文件一致。"""
import sys, io, os, json, time, subprocess, urllib.request, uuid, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = r"C:/Users/Cloak_Zeng/WorkBuddy/2026-10-06-14-57-26"
EXE = os.path.join(HERE, "mes_conv", "_newbuild.exe")
MES = os.path.join(HERE, "mes_conv")
PORT = int(os.environ.get("APP_PORT", "8761"))
BASE = "http://127.0.0.1:%d" % PORT
OUT = os.path.join(MES, "out")

PASS = FAIL = 0
def check(name, ok, extra=""):
    global PASS, FAIL
    if ok:
        PASS += 1; print("  ✓ %s" % name)
    else:
        FAIL += 1; print("  ✗ %s   %s" % (name, extra))

def multipart(fields, files):
    """fields: dict[str,str]; files: dict[str, list[(filename, bytes)]]"""
    b = "----wb" + uuid.uuid4().hex
    body = io.BytesIO()
    def w(s):
        body.write(s.encode("utf-8") if isinstance(s, str) else s)
    for k, v in fields.items():
        w("--%s\r\n" % b)
        w('Content-Disposition: form-data; name="%s"\r\n\r\n' % k)
        w(v); w("\r\n")
    for k, arr in files.items():
        for fn, data in arr:
            w("--%s\r\n" % b)
            w('Content-Disposition: form-data; name="%s"; filename="%s"\r\n' % (k, fn))
            w("Content-Type: application/octet-stream\r\n\r\n")
            w(data); w("\r\n")
    w("--%s--\r\n" % b)
    return body.getvalue(), "multipart/form-data; boundary=%s" % b

def post(path, fields, files, timeout=180):
    data, ct = multipart(fields, files)
    req = urllib.request.Request(BASE + path, data=data, headers={"Content-Type": ct})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def get(path, timeout=20):
    with urllib.request.urlopen(BASE + path, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def rd(p):
    return open(p, "rb").read()

# ---------- 启动服务 ----------
app = subprocess.Popen([EXE, "-nobrowser", "-nosingle", "-port", str(PORT)],
                       cwd=MES, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
up = False
for _ in range(40):
    time.sleep(0.4)
    try:
        get("/api/ping", 2); up = True; break
    except Exception:
        pass
print("服务已启动:", up)
if not up:
    app.kill(); sys.exit(1)

try:
    # 清掉旧产物，便于确认新文件
    for n in os.listdir(OUT) if os.path.isdir(OUT) else []:
        if n.startswith("MES物料档案_") or n.startswith("订单工单_"):
            try: os.remove(os.path.join(OUT, n))
            except Exception: pass

    S = r"D:/User/Cloak_Zeng/Download"
    tpl = rd(os.path.join(S, "物料档案导入模版 (3).xlsx"))
    full = rd(os.path.join(S, "20261006 匯入 577筆.xlsx"))
    p1, p2 = os.path.join(HERE, "_fmttest/拆_1.xlsx"), os.path.join(HERE, "_fmttest/拆_2.xlsx")

    print("\n【1】物料档案：两个源文件一次提交（应合并成一份、共 577 行）")
    r = post("/api/convert", {"lang": "zh"}, {
        "source": [("拆_1.xlsx", rd(p1)), ("拆_2.xlsx", rd(p2))],
        "template": [("tpl.xlsx", tpl)],
    })
    check("接口返回 ok", r.get("ok"), json.dumps(r, ensure_ascii=False)[:300])
    check("★ 合并后共 577 行", r.get("rows") == 577, "rows=%s" % r.get("rows"))
    batch_out = r.get("file")
    check("输出文件名带「批量」语义或时间戳", bool(batch_out), batch_out)
    logs = "\n".join(r.get("log", []))
    check("日志写明批量文件数", "2" in logs and ("批量" in logs or "批次" in logs), logs[:200])
    check("日志列出每个文件的行数", "288" in logs and "289" in logs, logs[:400])

    print("\n【2】同一份数据、单个文件提交（对照）")
    r2 = post("/api/convert", {"lang": "zh"}, {
        "source": [("full.xlsx", full)],
        "template": [("tpl.xlsx", tpl)],
    })
    check("接口返回 ok", r2.get("ok"), json.dumps(r2, ensure_ascii=False)[:200])
    check("★ 单文件也是 577 行", r2.get("rows") == 577, "rows=%s" % r2.get("rows"))

    print("\n【3】批量产物 vs 单文件产物：应逐格一致")
    sys.path.insert(0, HERE)
    import openpyxl

    def load_sheet(p):
        wb = openpyxl.load_workbook(p, data_only=True)
        sh = wb[wb.sheetnames[0]]
        rows = [["" if c is None else str(c).strip() for c in r] for r in sh.iter_rows(values_only=True)]
        wb.close(); return rows

    A = load_sheet(os.path.join(OUT, r2["file"]))
    B = load_sheet(os.path.join(OUT, r["file"]))
    diff = 0
    if len(A) != len(B):
        diff += 1
        print("     行数不同：单文件 %d / 批量 %d" % (len(A), len(B)))
    else:
        for i in range(len(A)):
            for j in range(max(len(A[i]), len(B[i]))):
                va = A[i][j] if j < len(A[i]) else ""
                vb = B[i][j] if j < len(B[i]) else ""
                if va != vb:
                    diff += 1
                    if diff <= 3:
                        print("     第%d行 第%d列 %r vs %r" % (i + 1, j + 1, va, vb))
    check("★ 批量合并结果与单文件逐格一致", diff == 0, "%d 处差异" % diff)

    print("\n【4】订单模块：两个订单文件一次提交")
    OS = r"D:/User/Cloak_Zeng/Code/python/project-001/code-file/015/文件"
    order = rd(os.path.join(OS, "訂單資料.xlsx"))
    work = rd(os.path.join(OS, "工單資料.xlsx"))
    r3 = post("/api/orders/convert", {"lang": "zh"}, {
        "order": [("訂單A.xlsx", order), ("訂單B.xlsx", order)],
        "work": [("工單資料.xlsx", work)],
    })
    check("接口返回 ok", r3.get("ok"), json.dumps(r3, ensure_ascii=False)[:300])
    t3 = r3.get("types") or {}
    check("★ 两个订单文件里的行都进来了（明细 6 行 = 3×2）",
          t3.get("销售订单明细") == 6, "types=%s" % t3)
    check("订单主表按单号去重后仍是 1 张", t3.get("销售订单主表") == 1, "types=%s" % t3)
    logs3 = "\n".join(r3.get("log", []))
    check("日志列出两个订单文件的行数说明", "· 訂單A.xlsx：3 行" in logs3 or "訂單A.xlsx" in logs3, logs3[:400])

    print("\n【5】批量说明（notes）")
    notes = "\n".join(r.get("notes", []))
    print("     notes:", notes[:400] or "(空)")
    check("★ 物料档案结果带批量说明", "批量" in notes or "合并" in notes, notes[:200])
    notes3 = "\n".join(r3.get("notes", []))
    check("★ 订单结果带批量说明", len(notes3) > 0, notes3[:200])
    print("     订单 notes:", notes3[:400] or "(空)")
    print("     订单 warnings:", "\n".join(r3.get("warnings", []))[:300] or "(空)")

finally:
    app.terminate()
    try: app.wait(timeout=8)
    except Exception: app.kill()

print("\n结果：%d 通过 / %d 失败" % (PASS, FAIL))
sys.exit(1 if FAIL else 0)
