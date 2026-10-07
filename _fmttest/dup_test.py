# -*- coding: utf-8 -*-
"""#36 重复行 / 重复编码 + Excel 错误值：端到端验证。

夹具全部由已有的 _fmttest/拆_1.xlsx（288 行真实数据）派生：
  dup_mat.xlsx  复制第 1 行 -> 物料编码重复 1 组
  err_mat.xlsx  把 MB002（映射到 物料名称）写成 #NAME? -> Excel 错误值
"""
import sys, io, os, re, json, time, subprocess, urllib.request, uuid
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

HERE = r"C:/Users/Cloak_Zeng/WorkBuddy/2026-10-06-14-57-26"
MES = os.path.join(HERE, "mes_conv")
EXE = os.path.join(MES, "_newbuild.exe")
PORT = int(os.environ.get("APP_PORT", "8763"))
BASE = "http://127.0.0.1:%d" % PORT
OUT = os.path.join(MES, "out")
FT = os.path.join(HERE, "_fmttest")
S = r"D:/User/Cloak_Zeng/Download"

PASS = FAIL = 0
def check(name, ok, extra=""):
    global PASS, FAIL
    if ok:
        PASS += 1; print("  OK  %s" % name)
    else:
        FAIL += 1; print("  NG  %s   %s" % (name, extra))

def multipart(fields, files):
    b = "----wb" + uuid.uuid4().hex
    body = io.BytesIO()
    def w(s):
        body.write(s.encode("utf-8") if isinstance(s, str) else s)
    for k, v in fields.items():
        w("--%s\r\n" % b)
        w('Content-Disposition: form-data; name="%s"\r\n\r\n' % k)
        w(v); w("\r\n")
    for k, arr in files.items():
        for fn, data in arr:
            w("--%s\r\n" % b)
            w('Content-Disposition: form-data; name="%s"; filename="%s"\r\n' % (k, fn))
            w("Content-Type: application/octet-stream\r\n\r\n")
            w(data); w("\r\n")
    w("--%s--\r\n" % b)
    return body.getvalue(), "multipart/form-data; boundary=%s" % b

def post(path, fields, files, timeout=180):
    data, ct = multipart(fields, files)
    req = urllib.request.Request(BASE + path, data=data, headers={"Content-Type": ct})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def get(path, timeout=20):
    with urllib.request.urlopen(BASE + path, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def rd(p): return open(p, "rb").read()

def join(seq): return "\n".join(seq or [])

# ---------- 造夹具 ----------
def make_fixtures():
    src = os.path.join(FT, "拆_1.xlsx")
    wb = openpyxl.load_workbook(src, data_only=True)
    sh = wb[wb.sheetnames[0]]
    rows = [list(r) for r in sh.iter_rows(values_only=True)]
    hdr, data = rows[0], rows[1:]
    wb.close()

    # a) 重复物料编码：复制第一条数据行
    p_dup = os.path.join(FT, "dup_mat.xlsx")
    w = openpyxl.Workbook(); s = w.active
    s.append(hdr)
    for r in data + [data[0]]:
        s.append(r)
    w.save(p_dup)

    # b) Excel 错误值：把 MB002（-> 物料名称）写成 #NAME?
    p_err = os.path.join(FT, "err_mat.xlsx")
    w = openpyxl.Workbook(); s = w.active
    s.append(hdr)
    for r in data:
        s.append(r)
    s.cell(row=2, column=2).value = "#NAME?"   # MB002
    w.save(p_err)
    return p_dup, p_err, len(data)

def open_out(name):
    p = os.path.join(OUT, name)
    wb = openpyxl.load_workbook(p, data_only=True)
    sh = wb[wb.sheetnames[0]]
    n = sh.max_row
    wb.close()
    return n

# ---------- 启动服务 ----------
p_dup, p_err, ndata = make_fixtures()
print("夹具：dup_mat.xlsx = %d 行（含 1 组重复）  err_mat.xlsx = %d 行（含 1 处 #NAME?）"
      % (ndata + 1, ndata))

app = subprocess.Popen([EXE, "-nobrowser", "-nosingle", "-port", str(PORT)],
                       cwd=MES, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
up = False
for _ in range(40):
    time.sleep(0.4)
    try:
        get("/api/ping", 2); up = True; break
    except Exception:
        pass
print("服务已启动:", up)
if not up:
    app.kill(); sys.exit(1)

try:
    tpl = rd(os.path.join(S, "物料档案导入模版 (3).xlsx"))
    OS = r"D:/User/Cloak_Zeng/Code/python/project-001/code-file/015/文件"
    order = rd(os.path.join(OS, "訂單資料.xlsx"))
    work = rd(os.path.join(OS, "工單資料.xlsx"))

    print("\n【1】物料档案：重复物料编码")
    r = post("/api/convert", {"lang": "zh"}, {
        "source": [("dup_mat.xlsx", rd(p_dup))],
        "template": [("tpl.xlsx", tpl)],
    })
    check("接口 ok", r.get("ok"), json.dumps(r, ensure_ascii=False)[:200])
    check("★ 行数 = %d（原 288 + 复制 1）" % (ndata + 1), r.get("rows") == ndata + 1, "rows=%s" % r.get("rows"))
    errs, notes = join(r.get("errors")), join(r.get("notes"))
    print("     notes:", notes.replace("\n", " | ")[:300])
    print("     errors:", errs.replace("\n", " | ")[:300])
    check("★ errors 指出重复", "重复" in errs, errs[:300])
    # 断言方式：从结果里正则取回被点名的编码，而不是写死具体料号（仓库公开，避免内联业务数据）
    m = re.search(r"重复[：:]\s*([0-9A-Za-z_\-]{6,})", errs)
    check("★ 重复提示点名了具体物料编码", bool(m), errs[:300])
    check("★ notes 汇总「发现 1 组重复」", "1 组重复" in notes, notes[:300])
    check("明细 issues 里也有重复条", any("重复" in (it.get("msg") or "") for it in (r.get("issues") or [])),
          str((r.get("issues") or [])[:2])[:300])
    check("转换未被打断（产物可打开）", open_out(r["file"]) == ndata + 2, "表格行数=%s" % open_out(r["file"]))

    print("\n【2】物料档案：Excel 错误值（MB002 = #NAME?）")
    r2 = post("/api/convert", {"lang": "zh"}, {
        "source": [("err_mat.xlsx", rd(p_err))],
        "template": [("tpl.xlsx", tpl)],
    })
    check("接口 ok", r2.get("ok"), json.dumps(r2, ensure_ascii=False)[:200])
    w2, n2 = join(r2.get("warnings")), join(r2.get("notes"))
    print("     warnings:", w2.replace("\n", " | ")[:300])
    print("     notes:", n2.replace("\n", " | ")[:200])
    check("★ warnings 指出 Excel 错误值", "Excel 错误值" in w2, w2[:300])
    check("★ 汇总写明 MB002 与格数", "MB002" in w2 and "1 处" in w2, w2[:300])
    check("★ issues 定位到该格", any("错误值" in (it.get("msg") or "") for it in (r2.get("issues") or [])),
          str((r2.get("issues") or [])[:2])[:200])
    check("★ notes 说明未发现重复", "未发现重复" in n2, n2[:200])

    print("\n【3】物料档案：干净数据（无重复）")
    r3 = post("/api/convert", {"lang": "zh"}, {
        "source": [("clean.xlsx", rd(os.path.join(FT, "拆_1.xlsx")))],
        "template": [("tpl.xlsx", tpl)],
    })
    n3 = join(r3.get("notes"))
    check("★ notes 写明未发现重复", "未发现重复" in n3, n3[:200])
    check("★ 重复检查带上列名", "物料编码" in n3, n3[:200])
    check("无重复时不进 errors", "重复" not in join(r3.get("errors")), join(r3.get("errors"))[:200])

    print("\n【4】订单：同一份订单源档提交两次（明细必然重复）")
    r4 = post("/api/orders/convert", {"lang": "zh"}, {
        "order": [("訂單A.xlsx", order), ("訂單B.xlsx", order)],
        "work": [("工單資料.xlsx", work)],
    })
    check("接口 ok", r4.get("ok"), json.dumps(r4, ensure_ascii=False)[:250])
    t4 = r4.get("types") or {}
    check("明细仍是 6 行", t4.get("销售订单明细") == 6, "types=%s" % t4)
    e4, n4 = join(r4.get("errors")), join(r4.get("notes"))
    print("     notes:", n4.replace("\n", " | ")[:400])
    print("     errors:", e4.replace("\n", " | ")[:400])
    check("★ 明细重复被抓到", "订单编号 + 产品编码" in e4, e4[:400])
    check("★ 汇总写明 3 组重复", "3 组重复" in n4, n4[:400])
    check("★ 工单表也做了重复检查（未发现重复）", "工单号" in n4, n4[:400])

    print("\n【5】订单：单份源档（明细无重复）")
    r5 = post("/api/orders/convert", {"lang": "zh"}, {
        "order": [("訂單A.xlsx", order)],
        "work": [("工單資料.xlsx", work)],
    })
    n5, e5 = join(r5.get("notes")), join(r5.get("errors"))
    print("     notes:", n5.replace("\n", " | ")[:400])
    check("★ 明细未发现重复", "订单编号 + 产品编码" in n5 and "未发现重复" in n5, n5[:400])
    check("★ 单份订单不进重复 errors", "重复" not in e5, e5[:300])

    print("\n【6】重复检查可配置：orders.json 覆盖后 3 组 -> 1 组")
    cfgp = os.path.join(MES, "orders.json")
    orig = open(cfgp, "rb").read()
    try:
        cfg = json.loads(orig.decode("utf-8"))
        cfg["duplicate_keys"] = {"order_detail": ["订单编号"], "work_order": []}
        open(cfgp, "wb").write(json.dumps(cfg, ensure_ascii=False, indent=2).encode("utf-8"))
        r6 = post("/api/orders/convert", {"lang": "zh"}, {
            "order": [("訂單A.xlsx", order), ("訂單B.xlsx", order)],
            "work": [("工單資料.xlsx", work)],
        })
        n6 = join(r6.get("notes"))
        print("     notes:", n6.replace("\n", " | ")[:400])
        check("★ 按订单编号单独成组：1 组重复", "1 组重复" in n6, n6[:400])
        check("★ 键名随配置变化（只有订单编号）",
              "重复检查（订单编号）" in n6, n6[:400])
        check("★ 工单配置为空数组 = 关闭检查", "工单号" not in n6, n6[:400])
    finally:
        open(cfgp, "wb").write(orig)

finally:
    app.terminate()
    try: app.wait(timeout=8)
    except Exception: app.kill()

print("\n结果：%d 通过 / %d 失败" % (PASS, FAIL))
sys.exit(1 if FAIL else 0)
