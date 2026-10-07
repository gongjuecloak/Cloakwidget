# -*- coding: utf-8 -*-
"""#35 模板列变化检测：基准记录 / 一致 / 新增列 / 缺列 / 换序 / 重置。"""
import sys, io, os, json, time, subprocess, urllib.request, uuid
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

HERE = r"C:/Users/Cloak_Zeng/WorkBuddy/2026-10-06-14-57-26"
MES = os.path.join(HERE, "mes_conv")
EXE = os.path.join(MES, "_newbuild.exe")
PORT = int(os.environ.get("APP_PORT", "8765"))
BASE = "http://127.0.0.1:%d" % PORT
FT = os.path.join(HERE, "_fmttest")
S = r"D:/User/Cloak_Zeng/Download"
BASELINE = os.path.join(MES, "tpl_baseline.json")

PASS = FAIL = 0
def check(name, ok, extra=""):
    global PASS, FAIL
    if ok: PASS += 1; print("  OK  %s" % name)
    else:  FAIL += 1; print("  NG  %s   %s" % (name, extra))

def multipart(fields, files):
    b = "----wb" + uuid.uuid4().hex
    body = io.BytesIO()
    def w(s): body.write(s.encode("utf-8") if isinstance(s, str) else s)
    for k, v in fields.items():
        w("--%s\r\n" % b); w('Content-Disposition: form-data; name="%s"\r\n\r\n' % k); w(v); w("\r\n")
    for k, arr in files.items():
        for fn, data in arr:
            w("--%s\r\n" % b)
            w('Content-Disposition: form-data; name="%s"; filename="%s"\r\n' % (k, fn))
            w("Content-Type: application/octet-stream\r\n\r\n"); w(data); w("\r\n")
    w("--%s--\r\n" % b)
    return body.getvalue(), "multipart/form-data; boundary=%s" % b

def post(path, fields, files, timeout=180):
    data, ct = multipart(fields, files)
    req = urllib.request.Request(BASE + path, data=data, headers={"Content-Type": ct})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def get(path, timeout=20):
    with urllib.request.urlopen(BASE + path, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def rd(p): return open(p, "rb").read()
def join(s): return "\n".join(s or [])

# ---------- 模板变体 ----------
tpl0 = rd(os.path.join(S, "物料档案导入模版 (3).xlsx"))

def load_tpl_headers(data, tmp):
    open(tmp, "wb").write(data)
    wb = openpyxl.load_workbook(tmp)
    sh = wb[wb.sheetnames[0]]
    hdr = [c.value for c in sh[1]]
    wb.close(); return hdr

def build_tpl(hdr, tmp):
    wb = openpyxl.Workbook(); sh = wb.active
    sh.append(hdr)
    wb.save(tmp)
    return rd(tmp)

tmp = os.path.join(FT, "_tpl_tmp.xlsx")
hdr0 = load_tpl_headers(tpl0, tmp)
ncol = len(hdr0)
print("模板列数：%d，前 5 列：%s" % (ncol, hdr0[:5]))

hdr_add = list(hdr0) + ["新列X"]
# 注意：模板列名带 "* " 必填前缀，比对要先剥掉前缀
def clean(h): return ("" if h is None else str(h).lstrip("*").strip())
hdr_drop = [h for h in hdr0 if clean(h) != "物料编码"]   # 删掉被映射的列
hdr_swap = list(hdr0); hdr_swap[0], hdr_swap[1] = hdr_swap[1], hdr_swap[0]
print("去掉物料编码后仍剩 %d 列" % len(hdr_drop))

tpl_add = build_tpl(hdr_add, tmp)
tpl_drop = build_tpl(hdr_drop, tmp)
tpl_swap = build_tpl(hdr_swap, tmp)

# ---------- 干净起点 ----------
if os.path.exists(BASELINE):
    os.remove(BASELINE)

app = subprocess.Popen([EXE, "-nobrowser", "-nosingle", "-port", str(PORT)],
                       cwd=MES, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
up = False
for _ in range(40):
    time.sleep(0.4)
    try: get("/api/ping", 2); up = True; break
    except Exception: pass
print("服务已启动:", up)
if not up: app.kill(); sys.exit(1)

def conv(tpl_bytes, name="tpl.xlsx"):
    return post("/api/convert", {"lang": "zh"}, {
        "source": [("拆_1.xlsx", rd(os.path.join(FT, "拆_1.xlsx")))],
        "template": [(name, tpl_bytes)],
    })

try:
    print("\n【1】首次转换：应记录基准")
    st0 = get("/api/tpl/status")
    check("重置后状态为「无基准」", st0.get("has") is False, json.dumps(st0, ensure_ascii=False))
    check("状态带中性说明", "尚未记录" in (st0.get("message") or ""), st0.get("message"))
    r1 = conv(tpl0)
    n1 = join(r1.get("notes"))
    print("     notes:", n1.replace("\n", " | ")[:300])
    check("★ 首次转换记下基准", ("%d 列" % ncol) in n1 and "基准" in n1, n1[:300])
    check("基准文件已落盘", os.path.exists(BASELINE))
    st1 = get("/api/tpl/status")
    check("★ 接口报告基准状态", st1.get("has") and st1.get("count") == ncol, json.dumps(st1, ensure_ascii=False))
    check("状态含记录时间", bool(st1.get("saved_at")), st1.get("saved_at"))

    print("\n【2】再次用同一个模板：应报「一致」")
    r2 = conv(tpl0)
    n2, w2 = join(r2.get("notes")), join(r2.get("warnings"))
    print("     notes:", n2.replace("\n", " | ")[:300])
    check("★ 提示与基准一致", "一致" in n2, n2[:300])
    check("一致时不产生模板告警", "模板" not in w2, w2[:200])

    print("\n【3】模板多了一列（未映射）")
    r3 = conv(tpl_add)
    n3, w3 = join(r3.get("notes")), join(r3.get("warnings"))
    print("     warnings:", w3.replace("\n", " | ")[:400])
    check("★ 报出新增列", "新增列「新列X」" in w3, w3[:400])
    check("★ 报出列数变化", ("由 %d 变为 %d" % (ncol, ncol + 1)) in w3, w3[:400])
    check("★ 列数变化文案与模板一致", (ncol + 1) == ncol + 1 and str(ncol + 1) in w3, w3[:200])

    print("\n【4】模板缺了被映射的列（物料编码）")
    r4 = conv(tpl_drop)
    w4 = join(r4.get("warnings"))
    print("     warnings:", w4.replace("\n", " | ")[:400])
    check("★ 报出缺失列", "少了列「物料编码」" in w4, w4[:400])

    print("\n【5】模板只是换了个序（集合相同）")
    r5 = conv(tpl_swap)
    w5 = join(r5.get("warnings"))
    print("     warnings:", w5.replace("\n", " | ")[:400])
    check("★ 报出列顺序变化", "顺序有变化" in w5, w5[:400])
    check("★ 顺序信息里带列名", ("物料编码" in w5) and ("物料名称" in w5), w5[:400])

    print("\n【6】清除基准后可重新记录")
    rr = post("/api/tpl/reset", {"lang": "zh"}, {})
    check("接口 ok", rr.get("ok"), json.dumps(rr, ensure_ascii=False))
    check("★ 返回清除说明", "清除模板基准" in (rr.get("message") or ""), rr.get("message"))
    check("基准文件已删除", not os.path.exists(BASELINE))
    st2 = get("/api/tpl/status")
    check("状态回到「无基准」", st2.get("has") is False, json.dumps(st2, ensure_ascii=False))
    r6 = conv(tpl_swap)
    n6 = join(r6.get("notes"))
    check("★ 下次转换按新模板重新记录", ("%d 列" % ncol) in n6 and "基准" in n6, n6[:300])
    r7 = conv(tpl_swap)
    check("★ 再用同一模板即「一致」", "一致" in join(r7.get("notes")), join(r7.get("notes"))[:200])

    print("\n【7】回归：模板检测不影响转换结果")
    check("首轮转换行数不变", r1.get("rows") == 288, "rows=%s" % r1.get("rows"))
    check("首轮转换仍无错误", not r1.get("errors"), str(r1.get("errors"))[:200])

finally:
    app.terminate()
    try: app.wait(timeout=8)
    except Exception: app.kill()

print("\n结果：%d 通过 / %d 失败" % (PASS, FAIL))
sys.exit(1 if FAIL else 0)
