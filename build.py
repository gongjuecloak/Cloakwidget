import openpyxl, sys

SRC = "D:/User/Cloak_Zeng/Download/20261006 匯入 577筆.xlsx"
TPL = "D:/User/Cloak_Zeng/Download/物料档案导入模版 (3).xlsx"
OUT = "D:/User/Cloak_Zeng/Download/物料档案导入_577筆.xlsx"

# --- target material_type dictionary (from 字典数据 (1).xls) ---
MAT_TYPE = {
    110: "成品", 120: "原材料", 130: "耗材", 140: "客供品", 150: "开发试作半成品",
    160: "半成品", 170: "费用类(五金/模具/开版)", 171: "收入类(模具/开版请款)",
    180: "管销类(事务性用品)", 190: "托工类(耗用内部材料)", 191: "托工类(纯托工)",
    210: "房屋建築", 211: "机器设备", 212: "运输设备", 213: "办公设备", 214: "其它固资",
    "R": "原材料", "S": "半成品", "F": "成品", "P": "产品", "C": "耗材", "T": "工具",
}
MAT_SOURCE = {"M": "自制", "P": "外购", "S": "委外"}

# --- load template headers ---
twb = openpyxl.load_workbook(TPL, data_only=True, read_only=True)
tws = twb[twb.sheetnames[0]]
theaders = [c.value for c in next(tws.iter_rows(min_row=1, max_row=1))]
twb.close()
print("template headers:", len(theaders))
# clean-name -> column index (headers may carry a '* ' required prefix)
hdr_idx = {h.replace("* ", "").strip(): i for i, h in enumerate(theaders)}

# --- load source ---
swb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
sws = swb[swb.sheetnames[0]]
srows = list(sws.iter_rows(values_only=True))
sheader = srows[0]
# build mb name -> col index
mbidx = {}
for i, h in enumerate(sheader):
    if h is not None:
        mbidx[str(h).strip()] = i
print("source MB count mapped:", len(mbidx))

def g(row, mb):
    j = mbidx.get(mb)
    if j is None or j >= len(row):
        return None
    v = row[j]
    return v

def norm_num(v):
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v

# --- build output ---
owb = openpyxl.Workbook()
ows = owb.active
ows.title = "物料档案"
ows.append(theaders)

n = 0
for ri, row in enumerate(srows[1:], start=1):
    if all((c is None or c == "") for c in row):
        continue
    out = [None] * len(theaders)
    # helper to set by clean header name
    def setc(name, val):
        out[hdr_idx[name]] = val

    # 物料编码
    code = g(row, "MB001")
    setc("物料编码", str(code).strip() if code is not None else None)
    # 物料名称
    setc("物料名称", g(row, "MB002"))
    # 物料规格
    setc("物料规格", g(row, "MB003"))
    # 物料分类 (material_catid default 0)
    setc("物料分类", 0)
    # 物料类型 (MB005 -> dict)
    mt = g(row, "MB005")
    if mt is not None:
        key = int(mt) if isinstance(mt, (int, float)) else str(mt).strip()
        setc("物料类型", MAT_TYPE.get(key, str(mt)))
    # 基本单位 (MB004 raw)
    setc("基本单位", g(row, "MB004"))
    # 物料来源 (MB025)
    ms = g(row, "MB025")
    if ms is not None:
        setc("物料来源", MAT_SOURCE.get(str(ms).strip(), str(ms)))
    # 默认供应商 (MB032)
    setc("默认供应商", g(row, "MB032"))
    # 图纸编号 (MB029)
    setc("图纸编号", g(row, "MB029"))
    # 物料长/宽/高 (MB093/094/095)
    setc("物料长度", norm_num(g(row, "MB093")))
    setc("物料宽度", norm_num(g(row, "MB094")))
    setc("物料高度", norm_num(g(row, "MB095")))
    # 批次管理 (MB022 N -> 否)
    bm = g(row, "MB022")
    setc("批次管理", "否" if str(bm).strip().upper() == "N" else ("是" if bm else "否"))
    # 序列管理 (per user: rules default -> 是)
    setc("序列管理", "是")
    # 保质期天数 (MB023)
    setc("保质期天数", norm_num(g(row, "MB023")))
    # 时效物料 / 危险品标识 / 关键物料 -> 否; 物料状态 -> 正常
    setc("时效物料", "否")
    setc("危险品标识", "否")
    setc("关键物料", "否")
    setc("物料状态", "正常")
    # ABC分类 (MB027)
    setc("ABC分类", norm_num(g(row, "MB027")))
    # flags
    setc("虚拟件标识", "否")
    setc("虚设件标识", "否")
    setc("可采购标识", "否")
    setc("可生产标识", "否")
    setc("可销售标识", "是")
    setc("可委外标识", "否")
    # 条码编号 (MB013)
    setc("条码编号", g(row, "MB013"))
    # 显示顺序
    setc("显示顺序", ri)
    ows.append(out)
    n += 1

owb.save(OUT)
swb.close()
print("WROTE rows:", n, "->", OUT)

# --- validation summary ---
print("\n=== 抽样校验 (按 MB005 / MB025 取值) ===")
# re-open to verify
vw = openpyxl.load_workbook(OUT, data_only=True)
vs = vw.active
vheaders = [c.value for c in next(vs.iter_rows(min_row=1, max_row=1))]
vrows = list(vs.iter_rows(values_only=True))[1:]
print("output data rows:", len(vrows))
def col(name): return hdr_idx[name]
checks = {"物料类型": {}, "物料来源": {}, "序列管理": set(), "可销售标识": set(), "物料状态": set(), "批次管理": set()}
for r in vrows:
    checks["物料类型"][r[col("物料类型")]] = checks["物料类型"].get(r[col("物料类型")],0)+1
    checks["物料来源"][r[col("物料来源")]] = checks["物料来源"].get(r[col("物料来源")],0)+1
    checks["序列管理"].add(r[col("序列管理")])
    checks["可销售标识"].add(r[col("可销售标识")])
    checks["物料状态"].add(r[col("物料状态")])
    checks["批次管理"].add(r[col("批次管理")])
for k,v in checks.items():
    print(f"  {k}: {v}")
# show one row where 物料类型=开发试作半成品
for r in vrows:
    if r[col("物料类型")] == "开发试作半成品":
        print("  SAMPLE 开发试作半成品:", r[col("物料编码")], "|", r[col("物料名称")][:30] if r[col("物料名称")] else None, "| 来源=", r[col("物料来源")])
        break
vw.close()
