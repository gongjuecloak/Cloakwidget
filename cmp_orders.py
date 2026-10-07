# -*- coding: utf-8 -*-
"""把 Go 版订单/工单输出与 Python 版基准逐格比对。"""
import sys, io, os
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
GOLD = r"C:/Users/Cloak_Zeng/WorkBuddy/2026-10-06-14-57-26/_golden"


def norm(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    return s


def load(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return None
    sh = wb[sheet]
    rows = [[norm(c) for c in r] for r in sh.iter_rows(values_only=True)]
    wb.close()
    return rows


def numeq(va, vb):
    """两侧都是数字且相对误差 < 1e-6 视为同一数值（用于识别 float32 精度噪声）。"""
    try:
        fa, fb = float(va), float(vb)
    except (TypeError, ValueError):
        return False
    if fa == fb:
        return True
    scale = max(abs(fa), abs(fb), 1e-9)
    return abs(fa - fb) / scale < 1e-6


def cmp(name, gold_path, gold_sheet, go_path, go_sheet):
    a = load(gold_path, gold_sheet)
    b = load(go_path, go_sheet)
    if a is None:
        print(f"✗ {name}: Python 基准里没有 sheet「{gold_sheet}」")
        return False
    if b is None:
        print(f"✗ {name}: Go 输出里没有 sheet「{go_sheet}」")
        return False
    if len(a) != len(b):
        print(f"✗ {name}: 行数不同 Python={len(a)} Go={len(b)}")
    ha, hb = a[0] if a else [], b[0] if b else []
    hdr_ok = ha == hb
    if not hdr_ok:
        print(f"✗ {name}: 表头不同")
        print("   Python:", ha)
        print("   Go    :", hb)
    diffs = []
    pres = []
    n = max(len(a), len(b))
    for i in range(1, n):
        ra = a[i] if i < len(a) else []
        rb = b[i] if i < len(b) else []
        w = max(len(ra), len(rb))
        for j in range(w):
            va = ra[j] if j < len(ra) else ""
            vb = rb[j] if j < len(rb) else ""
            if va != vb:
                col = ha[j] if j < len(ha) else f"col{j}"
                if numeq(va, vb):
                    pres.append((i + 1, col, va, vb))
                else:
                    diffs.append((i + 1, col, va, vb))
    if pres:
        print(f"  ~ {name}: {len(pres)} 处「精度噪声」（数值等价，Python 侧 float32 降精度所致）")
        for r, c, va, vb in pres[:6]:
            print(f"     第{r}行 [{c}] Python={va!r}  Go={vb!r}")
    if diffs:
        print(f"✗ {name}: {len(diffs)} 处差异（显示前 12）")
        for r, c, va, vb in diffs[:12]:
            print(f"     第{r}行 [{c}] Python={va!r}  Go={vb!r}")
    else:
        print(f"✓ {name}: {len(a)-1} 行 × {len(ha)} 列  零实质差异")
    return (not diffs) and hdr_ok


ok = True
ok &= cmp("销售订单主表", os.path.join(GOLD, "py_订单.xlsx"), "销售订单主表",
          os.path.join(GOLD, "go_订单工单.xlsx"), "销售订单主表")
ok &= cmp("销售订单明细", os.path.join(GOLD, "py_订单.xlsx"), "销售订单明细",
          os.path.join(GOLD, "go_订单工单.xlsx"), "销售订单明细")
ok &= cmp("工单", os.path.join(GOLD, "py_工单.xlsx"), "工单",
          os.path.join(GOLD, "go_订单工单.xlsx"), "工单")
print()
print("结果：", "全部一致 ✓" if ok else "存在差异 ✗")
