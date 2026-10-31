# -*- coding: utf-8 -*-
"""验证单位归一：PCS 直通 / pcs→PCS / PCSQ→PCS / 捲→卷 / KG→千克 / M→米；SET 直通。"""
import sys, io, os, subprocess
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
MES = os.path.join(HERE, 'mes_conv')
SRC = os.environ.get('MES_SRC_FILE', os.path.join(HERE, '_fixtures', 'materials.xlsx'))
TPL = os.environ.get('MES_TPL_FILE', os.path.join(HERE, '_fixtures', 'materials_template.xlsx'))
EXE = os.environ.get('MES_EXE', os.path.join(MES, '物料档案转换工具.exe'))
for label, path in [('源文件', SRC), ('模板', TPL), ('程序', EXE)]:
    if not os.path.isfile(path):
        raise SystemExit(f'{label}不存在：{path}（可用 MES_SRC_FILE / MES_TPL_FILE / MES_EXE 指定）')

wb = openpyxl.load_workbook(SRC)
sh = wb[wb.sheetnames[0]]
rows = list(sh.iter_rows(values_only=True))
wb.close()
hdr = list(rows[0])
data = rows[1:11]

vals = ['PCS', 'pcs', 'PCSQ', '捲', '卷', 'KG', 'kg', 'M', 'm', 'SET']
nb = openpyxl.Workbook(); ns = nb.active; ns.title = 't'
ns.append(hdr)
for i, r in enumerate(data):
    r = list(r); r[3] = vals[i]; ns.append(r)
test_src = os.path.join(HERE, 'test_unit_norm.xlsx')
nb.save(test_src)
print('测试源已生成:', test_src)

test_out = os.path.join(HERE, 'test_unit_norm_out.xlsx')
if os.path.exists(test_out):
    os.remove(test_out)
r = subprocess.run([EXE, '-cli', '-lang', 'zh', '-src', test_src, '-tpl', TPL, '-out', test_out],
                   capture_output=True)
print('exit =', r.returncode)
out = (r.stdout.decode('utf-8', 'replace') + r.stderr.decode('utf-8', 'replace')).strip()
print('--- CLI 输出 ---')
print(out[-900:])
print('--- 检查输出 ---')

ob = openpyxl.load_workbook(test_out, data_only=True)
osh = ob[ob.sheetnames[0]]
oit = osh.iter_rows(values_only=True)
oh = [('' if c is None else str(c)) for c in next(oit)]
ci = [i for i, h in enumerate(oh) if '基本单位' in h][0]
got = [('' if row[ci] is None else str(row[ci])) for row in oit]
ob.close()

exp = ['PCS', 'PCS', 'PCS', '卷', '卷', '千克', '千克', '米', '米', 'SET']
print()
print('%-8s %-8s %-8s %s' % ('源值', '期望', '实际', '结果'))
ok = True
for a, b, c in zip(vals, exp, got):
    good = (b == c)
    ok = ok and good
    print('%-8s %-8s %-8s %s' % (a, b, c, 'OK' if good else 'FAIL'))
print()
print('=== 单位归一全部通过 ===' if ok else '=== 存在失败 ===')
