# -*- coding: utf-8 -*-
"""逐格比对 工具输出 vs 参考文件 物料档案导入_577筆.xlsx"""
import sys, io, os
from collections import Counter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
A = os.environ.get("MES_VERIFY_OUT", os.path.join(ROOT, "verify_out.xlsx"))
B = os.environ.get("MES_REFERENCE_OUT", os.path.join(ROOT, "_fixtures", "materials_reference.xlsx"))
for label, path in (("待验证输出", A), ("参考文件", B)):
    if not os.path.isfile(path):
        raise SystemExit(f"{label}不存在：{path}（可用 MES_VERIFY_OUT / MES_REFERENCE_OUT 指定）")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, str):
        s = v.strip()
        # 纯数字字符串归一
        try:
            f = float(s)
            if f == int(f):
                return str(int(f))
        except ValueError:
            pass
        return s
    return str(v).strip()


def load(p):
    # 注意：excelize 产物在 openpyxl read_only 模式下 dimension 解析异常，故用普通模式
    wb = openpyxl.load_workbook(p, data_only=True)
    sh = wb[wb.sheetnames[0]]
    rows = [[norm(c) for c in r] for r in sh.iter_rows(values_only=True)]
    wb.close()
    return rows


a, b = load(A), load(B)
ha, hb = a[0], b[0]
print("输出 %d 行 x %d 列 / 参考 %d 行 x %d 列" % (len(a) - 1, len(ha), len(b) - 1, len(hb)))

# 表头比对
hm = [(i, ha[i], hb[i]) for i in range(min(len(ha), len(hb))) if ha[i] != hb[i]]
print("表头差异:", hm if hm else "无 ✓")

# 数据比对（按行号对齐）
diff = Counter()
samples = {}
nrows = min(len(a), len(b)) - 1
for r in range(1, nrows + 1):
    ra, rb = a[r], b[r]
    for c in range(min(len(ra), len(rb))):
        if ra[c] != rb[c]:
            key = ha[c] if c < len(ha) else "col%d" % c
            diff[key] += 1
            if key not in samples:
                samples[key] = (r, ra[c], rb[c], ra[0])

print("\n总计差异单元格:", sum(diff.values()))
for k, v in diff.most_common():
    r, va, vb, code = samples[k]
    print("  [%s] x%d  例(行%d 物料%s): 输出=%r 参考=%r" % (k, v, r, code, va, vb))

# 基本单位专项
if "基本单位" in ha:
    ci = ha.index("基本单位")
    ca = Counter(ra[ci] for ra in a[1:])
    cb = Counter(rb[ci] for rb in b[1:])
    print("\n基本单位 输出分布:", dict(ca))
    print("基本单位 参考分布:", dict(cb))
    print("完全一致 ✓" if ca == cb else "★ 有差异")
